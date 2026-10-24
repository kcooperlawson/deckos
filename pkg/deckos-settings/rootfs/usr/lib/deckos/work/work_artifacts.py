"""Local document and image tools, executed by the optional artifact runtime."""
import json
import sys
from pathlib import Path


def create_document(data):
    path = Path(data['path'])
    title = str(data.get('title') or path.stem)
    sections = data.get('sections', [])
    rows = data.get('rows', [])
    suffix = path.suffix.lower()
    if suffix == '.docx':
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        doc = Document()
        doc.sections[0].header.paragraphs[0].text = 'DECKOS WORK'
        normal = doc.styles['Normal']
        normal.font.name, normal.font.size = 'Calibri', Pt(11)
        normal.font.color.rgb = RGBColor.from_string('263449')
        doc.add_heading(title, 0)
        for section in sections:
            if section.get('heading'):
                doc.add_heading(str(section['heading']), 1)
            for paragraph in str(section.get('text','')).split('\n'):
                doc.add_paragraph(paragraph)
        if rows:
            table = doc.add_table(rows=0, cols=max(map(len,rows)))
            table.style = 'Light Shading Accent 1'
            for row in rows:
                cells = table.add_row().cells
                for i, value in enumerate(row):
                    cells[i].text = str(value)
        doc.save(path)
    elif suffix == '.pdf':
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.enums import TA_LEFT
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from xml.sax.saxutils import escape
        styles = getSampleStyleSheet()
        story = [Paragraph(escape(title), styles['Title']), Spacer(1, 20)]
        for section in sections:
            if section.get('heading'):
                story.append(Paragraph(escape(str(section['heading'])), styles['Heading1']))
            for paragraph in str(section.get('text','')).split('\n'):
                story.extend([Paragraph(escape(paragraph), styles['BodyText']), Spacer(1,8)])
        if rows:
            table = Table([[Paragraph(escape(str(c)),styles['BodyText']) for c in r] for r in rows], repeatRows=1)
            table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#E7EFFC')),
                                       ('VALIGN',(0,0),(-1,-1),'TOP'),('GRID',(0,0),(-1,-1),.4,colors.lightgrey)]))
            story.append(table)
        SimpleDocTemplate(str(path), title=title, author='DeckOS Work').build(story)
    elif suffix == '.pptx':
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        for i, section in enumerate([{'heading':title,'text':'Created locally with DeckOS Work'}, *sections]):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor.from_string('0F1115')
            box = slide.shapes.add_textbox(Inches(.8), Inches(.7), Inches(11.7), Inches(1.3))
            p = box.text_frame.paragraphs[0]
            p.text = str(section.get('heading') or title)
            p.font.size, p.font.bold, p.font.color.rgb = Pt(34), True, RGBColor.from_string('E8EAF0')
            body = slide.shapes.add_textbox(Inches(.85), Inches(2.1), Inches(11.5), Inches(4.5)).text_frame
            body.word_wrap = True
            for j, text in enumerate(str(section.get('text','')).split('\n')):
                p = body.paragraphs[0] if j == 0 else body.add_paragraph()
                p.text = text
                p.font.size, p.font.color.rgb = Pt(22), RGBColor.from_string('C5CCDA')
        prs.save(path)
    elif suffix == '.xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        book = Workbook()
        sheet = book.active
        sheet.title = 'Sheet 1'
        for row in rows or [['Title','Details'],[title,'\n'.join(str(s.get('text','')) for s in sections)]]:
            # Formulas must be an explicit formula object, never an accidental string from imported text.
            sheet.append([cell['formula'] if isinstance(cell,dict) and 'formula' in cell else cell for cell in row])
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='1C2E4B')
            cell.font = Font(color='FFFFFF', bold=True)
        for column in sheet.columns:
            letter = get_column_letter(column[0].column)
            sheet.column_dimensions[letter].width = min(55,max(16,max(len(str(c.value or '')) for c in column)+3))
            for cell in column:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if isinstance(cell.value,str) and cell.value.startswith('='):
                    # The API distinguishes formulas from literal strings below.
                    original = (rows or [])[cell.row-1][cell.column-1] if rows else None
                    if not isinstance(original,dict):
                        cell.data_type = 's'
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        book.save(path)
    else:
        raise ValueError('Use .docx, .pdf, .pptx or .xlsx for structured documents; write_file handles text formats.')
    return {'path':str(path),'bytes':path.stat().st_size}


def extract(path):
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(path)
        return '\n\n'.join((p.extract_text() or '') for p in reader.pages[:100])[:80000]
    if suffix == '.docx':
        from docx import Document
        doc = Document(path)
        return '\n'.join([p.text for p in doc.paragraphs]+[' | '.join(c.text for c in r.cells) for t in doc.tables for r in t.rows])[:80000]
    if suffix == '.pptx':
        from pptx import Presentation
        return '\n'.join(s.text for slide in Presentation(path).slides for s in slide.shapes if s.has_text_frame)[:80000]
    if suffix == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(path,read_only=True,data_only=False)
        return '\n'.join(f'{sheet.title}: '+ ' | '.join(str('' if v is None else v) for v in row)
                         for sheet in book for row in sheet.iter_rows(max_row=1000,values_only=True))[:80000]
    return path.read_text(errors='replace')[:80000]


def generate_image(data):
    import torch
    from diffusers import AutoPipelineForImage2Image, AutoPipelineForText2Image
    from PIL import Image
    model = data.get('model') or 'stabilityai/stable-diffusion-xl-base-1.0'
    device = 'cuda' if torch.cuda.is_available() else 'cpu'
    dtype = torch.float16 if device == 'cuda' else torch.float32
    kwargs = dict(torch_dtype=dtype, use_safetensors=True)
    if device == 'cuda':
        kwargs['variant'] = 'fp16'
    cls = AutoPipelineForImage2Image if data.get('input_image') else AutoPipelineForText2Image
    print('Loading local image weights…',flush=True)
    pipe = cls.from_pretrained(model, **kwargs)
    if device == 'cuda':
        # Keep text encoders off the GPU while the denoiser runs. Sliced attention also
        # avoids the multi-GB math-SDPA allocation seen on RDNA4 with ROCm.
        pipe.enable_model_cpu_offload()
        pipe.enable_attention_slicing('auto')
        pipe.enable_vae_slicing()
        pipe.enable_vae_tiling()
    else:
        pipe.to(device)
    print('Image weights loaded. Starting local inference…',flush=True)
    width, height = int(data.get('width',1024)), int(data.get('height',1024))
    width, height = max(256,min(1536,width))//64*64, max(256,min(1536,height))//64*64
    args = dict(prompt=data['prompt'], num_inference_steps=max(1,min(60,int(data.get('steps',20)))),
                guidance_scale=0.0 if 'turbo' in model.lower() else 7.0,
                generator=torch.Generator(device='cpu').manual_seed(int(data.get('seed',42))))
    if data.get('input_image'):
        args.update(image=Image.open(data['input_image']).convert('RGB').resize((width,height)), strength=.7)
    else:
        args.update(width=width,height=height)
    image = pipe(**args).images[0]
    path = Path(data['path'])
    from PIL.PngImagePlugin import PngInfo
    meta = PngInfo()
    meta.add_text('prompt',data['prompt'])
    meta.add_text('model',model)
    meta.add_text('seed',str(data.get('seed',42)))
    image.save(path,pnginfo=meta)
    return {'path':str(path),'device':torch.cuda.get_device_name(0) if device=='cuda' else 'CPU','model':model}


if __name__ == '__main__':
    try:
        request = json.load(sys.stdin)
        action = request.pop('action')
        if action == 'document':
            result = create_document(request)
        elif action == 'image':
            result = generate_image(request)
        elif action == 'extract':
            result = {'text':extract(request['path'])}
        else:
            raise ValueError('Unknown artifact action')
        print(json.dumps(result),flush=True)
    except Exception as error:
        print(json.dumps({'error':str(error)}),flush=True)
        sys.exit(1)
